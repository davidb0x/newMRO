import re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import pdfplumber
from pathlib import Path

def clean_name(text):
    if not text:
        return text
    replacements = {
        'ą': 'a', 'ć': 'c', 'ę': 'e', 'ł': 'l', 'ń': 'n', 'ó': 'o', 'ś': 's', 'ź': 'z', 'ż': 'z',
        'Ą': 'A', 'Ć': 'C', 'Ę': 'E', 'Ł': 'L', 'Ń': 'N', 'Ó': 'O', 'Ś': 'S', 'Ź': 'Z', 'Ż': 'Z',
        '*': ''
    }
    for old_char, new_char in replacements.items():
        text = text.replace(old_char, new_char)
    return text


def normalize_whitespace(value):
    if value is None:
        return value
    text = str(value)
    return re.sub(r"\s+", " ", text).strip()

def extract_products_from_pdf(pdf_path, verbose=False):
    """
    Ekstraktuje dane produktów z pliku PDF Festo
    """
    products = []

    def log(message):
        if verbose:
            print(message, flush=True)
    
    try:
        with pdfplumber.open(pdf_path) as pdf:
            lines = []

            log(f"  -> Otwieram PDF: {pdf_path}")
            log(f"  -> Liczba stron: {len(pdf.pages)}")

            # Zbierz cały tekst z PDF
            raw_text = ""
            for page_index, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                log(f"  -> Strona {page_index}: {len(page_text)} znakow tekstu")
                raw_text += page_text + "\n"

            # Usuń bloki HTML
            raw_text = re.sub(r'<html.*?</html\s*>', '', raw_text, flags=re.DOTALL | re.IGNORECASE)
            
            lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
            full_text = "\n".join(lines)

            # Dzielimy tekst na sekcje według nagłówka pozycji
            sections = re.split(r'(Pozycja nr \d+ / Państwa pozycja nr \d+)', full_text)
            log(f"  -> Linii tekstu: {len(lines)}")
            log(f"  -> Sekcji pozycji: {len(sections) // 2}")

            for i in range(1, len(sections), 2):
                if i + 1 >= len(sections):
                    break

                position_header = sections[i]
                section_content = sections[i + 1]

                log(f"  -> Sekcja {i // 2 + 1}: {position_header.strip()[:120]}")
                
                # Material description (np. "Elektrozawór") is at the end of the previous section
                prev_section = sections[i - 1].strip()
                material_desc = ""
                if prev_section:
                    prev_lines = prev_section.split('\n')
                    if prev_lines:
                        last_line = prev_lines[-1].strip()
                        if not re.match(r'^(\d+,\d+|\d+|SZT|EUR)$', last_line):
                            material_desc = last_line
                if material_desc:
                    log(f"     Opis materialu: {material_desc}")

                # Wyciągnij numer pozycji
                position_match = re.search(r'Pozycja nr (\d+)', position_header)
                if not position_match:
                    log("     Brak numeru pozycji - pomijam")
                    continue

                position_nr = position_match.group(1)
                log(f"     Pozycja nr: {position_nr}")

                # KROK 1: Znajdź nazwę produktu (ciąg wielkich liter, cyfr i myślników)
                name_pattern = r'\b([A-Z]{2,}[-A-Z0-9]+(?:-[A-Z0-9]+)*)\b'
                name_matches = re.finditer(name_pattern, section_content[:500])

                product_name = None
                for match in name_matches:
                    candidate = match.group(1)
                    if len(candidate) > 5 and candidate not in ["EUR", "SZT", "ISO"]:
                        product_name = candidate
                        break

                if not product_name:
                    log("     Brak nazwy produktu - pomijam")
                    continue
                
                typ_produktu = product_name
                
                if material_desc:
                    product_name = f"{material_desc} {product_name}"
                    
                product_name = clean_name(product_name)

                log(f"     Typ: {typ_produktu}")
                log(f"     Nazwa: {product_name}")

                # KROK 2: Numer produktu i cena jednostkowa
                # Układ z surowego tekstu: "575525 63,22 1" lub "1205861 1,98 20"
                # potem "SZT" i wartosc netto
                product_number = "N/A"
                price = "N/A"
                price_source = ""

                table_pattern = r'(\d{6,7})\s+(\d+[,\.]\d{2})\s+\d+\b'
                table_match = re.search(table_pattern, section_content)
                if table_match:
                    product_number = table_match.group(1)
                    price = table_match.group(2).replace('.', ',')
                    price_source = "tabela"
                else:
                    number_match = re.search(r'(\d{6,7})\s+SZT', section_content)
                    if number_match:
                        product_number = number_match.group(1)

                    if product_number != "N/A":
                        after_number = section_content.split(product_number, 1)[-1]
                        price_match = re.search(r'(\d+[,\.]\d{2})', after_number)
                        if price_match:
                            price = price_match.group(1).replace('.', ',')
                            price_source = "tekst po numerze"
                    else:
                        price_match = re.search(r'(\d+[,\.]\d{2})', section_content)
                        if price_match:
                            price = price_match.group(1).replace('.', ',')
                            price_source = "tekst"

                log(f"     Numer produktu: {product_number}")
                log(f"     Cena: {price}{' (z ' + price_source + ')' if price_source else ''}")

                # KROK 3: Dni dostawy (max z zakresu + 7)
                delivery_days = "N/A"
                delivery_match = re.search(
                    r'Dostawa\s+(\d+)\s*-\s*(\d+)\s+dni\s+robocz',
                    section_content,
                    re.IGNORECASE,
                )
                if delivery_match:
                    max_days = int(delivery_match.group(2))
                    delivery_days = str(max_days + 7)
                    log(
                        "     Dostawa: "
                        f"{delivery_match.group(1)}-{delivery_match.group(2)} dni roboczych"
                        f" => {delivery_days}"
                    )
                else:
                    log("     Dostawa: brak informacji")

                products.append({
                    'pozycja': position_nr,
                    'nazwa': product_name,
                    'typ': typ_produktu,
                    'nr_produktu': product_number,
                    'cena': price,
                    'dni_dostawy': delivery_days
                })

            log(f"  -> Wyekstrahowano {len(products)} pozycji z pliku")
    
    except Exception as e:
        print(f"❌ Błąd podczas przetwarzania {pdf_path}: {str(e)}")
        return []
    
    return products

def extract_products_alternative_pattern(text):
    """
    Alternatywny wzorzec ekstrakcji - bardziej elastyczny
    """
    products = []
    
    # Dzielimy tekst na sekcje według "Pozycja nr"
    sections = re.split(r'Pozycja nr (\d+)', text)
    
    for i in range(1, len(sections), 2):
        if i + 1 < len(sections):
            position_nr = sections[i]
            section_text = sections[i + 1]
            
            # Szukamy nazwy produktu (ciąg wielkich liter, cyfr i myślników)
            name_match = re.search(r'\b([A-Z]{2,}[-A-Z0-9]+)\b', section_text)
            
            # Szukamy numeru produktu (6-7 cyfr przed "SZT")
            number_match = re.search(r'(\d{6,7})\s+SZT', section_text)
            
            # Szukamy ceny (format: XX,XX lub XXX,XX)
            price_match = re.search(r'(\d+[,\.]\d{2})\s+\d+[,\.]\d{2}', section_text)
            
            if name_match:
                products.append({
                    'pozycja': position_nr,
                    'nazwa': clean_name(name_match.group(1)),
                    'typ': name_match.group(1),
                    'nr_produktu': number_match.group(1) if number_match else "N/A",
                    'cena': price_match.group(1).replace('.', ',') if price_match else "N/A"
                })
    
    return products

def create_excel_from_products(products, output_filename='Oferta_Festo.xlsx'):
    """
    Tworzy plik Excel z wyekstrahowanych danych produktów
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet()
    ws.title = "Oferta Festo"
    
    # 🎨 Stylowanie nagłówków
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    
    # Nagłówki
    headers = ['Nazwa', 'Typ', 'Numer produktu', 'Wart. jednostkowa EUR', 'Dni dostawy', 'Nazwa + Numer']
    ws.append(headers)
    
    # Formatowanie nagłówków
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Dodawanie danych
    for product in products:
        nazwa = normalize_whitespace(product['nazwa'])
        typ_produktu = normalize_whitespace(product.get('typ', ''))
        nr_produktu = normalize_whitespace(product['nr_produktu'])
        cena = normalize_whitespace(product.get('cena', ''))
        dni_dostawy = normalize_whitespace(product.get('dni_dostawy', 'N/A'))
        nazwa_numer = normalize_whitespace(f"{nazwa} {nr_produktu}".strip())
        ws.append([
            nazwa,
            typ_produktu,
            nr_produktu,
            cena,
            dni_dostawy,
            nazwa_numer,
        ])
    
    # Dostosowanie szerokości kolumn
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 25
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 45
    
    # Wyrównanie danych
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[0].alignment = Alignment(horizontal='left', vertical='center')
        row[1].alignment = Alignment(horizontal='left', vertical='center')
        row[2].alignment = Alignment(horizontal='center', vertical='center')
        row[3].alignment = Alignment(horizontal='right', vertical='center')
        row[4].alignment = Alignment(horizontal='center', vertical='center')
        row[5].alignment = Alignment(horizontal='left', vertical='center')
    
    # Zapisanie pliku
    wb.save(output_filename)
    print(f"✅ Plik {output_filename} został utworzony!")
    print(f"📊 Dodano {len(products)} produktów")

def process_all_pdfs_in_folder(
    folder_path='.',
    output_filename='Oferta_Festo.xlsx',
    verbose=True,
):
    """
    Przetwarza wszystkie pliki PDF w folderze i tworzy jeden plik Excel
    """
    all_products = []
    pdf_candidates = list(Path(folder_path).glob('*.pdf')) + list(Path(folder_path).glob('*.PDF'))
    pdf_files = []
    seen_paths = set()
    for pdf_file in pdf_candidates:
        try:
            normalized = str(pdf_file.resolve()).lower()
        except Exception:
            normalized = str(pdf_file).lower()
        if normalized in seen_paths:
            continue
        seen_paths.add(normalized)
        pdf_files.append(pdf_file)
    
    if not pdf_files:
        print("⚠️ Nie znaleziono plików PDF w folderze!")
        return
    
    print(f"🔍 Znaleziono {len(pdf_files)} plików PDF")
    if verbose:
        for pdf_file in pdf_files:
            print(f"  - {pdf_file}", flush=True)
    
    for pdf_file in pdf_files:
        print(f"\n📄 Przetwarzam: {pdf_file.name}")
        products = extract_products_from_pdf(pdf_file, verbose=verbose)
        
        if products:
            print(f" ✅ Wyekstrahowano {len(products)} produktów")
            all_products.extend(products)
        else:
            print(f" ⚠️ Nie znaleziono produktów w tym pliku")
    
    if all_products:
        # Usuń duplikaty (na podstawie nazwy produktu)
        unique_products = []
        seen_names = set()
        
        for product in all_products:
            if product['nazwa'] not in seen_names:
                unique_products.append(product)
                seen_names.add(product['nazwa'])
        
        print(f"\n📋 Łącznie unikalnych produktów: {len(unique_products)}")
        create_excel_from_products(unique_products, output_filename)
    else:
        print("\n❌ Nie udało się wyekstrahować żadnych produktów!")

def main():
    """
    Główna funkcja programu
    """
    print("🚀 Automatyczna ekstrakcja danych z plików PDF Festo\n")

    user_path = input(
        "Podaj sciezke do folderu z PDF-ami (Enter = biezacy katalog): "
    ).strip()

    target_path = Path(user_path) if user_path else Path(".")
    output_file = "Oferta_Festo.xlsx"

    if target_path.is_file():
        target_path = target_path.parent

    if target_path.is_dir():
        process_all_pdfs_in_folder(str(target_path), output_file)
    else:
        print("❌ Podana sciezka nie istnieje.")

    print("\n✨ Zakonczono!")

if __name__ == "__main__":
    main()